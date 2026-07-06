"""Complete the Soft Sourdough Master Tracker by filling blank cells from Google Drive."""
import sys
sys.path.insert(0, r"D:\TMB 2.0")
import warnings; warnings.filterwarnings("ignore")
import pandas as pd
from datetime import date, timedelta
from openpyxl import load_workbook
from services.gdrive import get_service, list_files_in_folder, download_file
from services.data_processor import extract_date_from_filename, clean_product_name
from config import TMB_SALES_FOLDER_ID

TRACKER_PATH = r"D:\Pastry Performace\Master Tracker.xlsx"

PRODUCTS = ["House Sourdough", "XL House Sourdough", "Tinned House Sourdough", "Soft Sourdough"]
DAY_REV_COLS = ["B","D","F","H","J","L","N"]
DAY_QTY_COLS = ["C","E","G","I","K","M","O"]

WEEKS = [
    (date(2026,3,2),  6),
    (date(2026,3,9),  15),
    (date(2026,3,16), 24),
    (date(2026,3,23), 33),
    (date(2026,3,30), 42),
    (date(2026,4,6),  51),
    (date(2026,4,13), 60),
    (date(2026,4,20), 69),
    (date(2026,4,27), 78),
    (date(2026,5,4),  87),
    (date(2026,5,11), 96),
    (date(2026,5,18), 105),
]


def categorise(desc: str) -> str | None:
    """Map a cleaned product description to one of the 4 tracker buckets."""
    d = desc.upper().strip()
    if "SLICED HOUSE SOURDOUGH" in d:
        return "House Sourdough"
    if ("HOUSE SOURDOUGH XL" in d) or ("XL HOUSE SOURDOUGH" in d) or ("XL HSE SOURDOUGH" in d):
        return "XL House Sourdough"
    if "TINNED HOUSE SOURDOUGH" in d or d == "TIN HOUSE SOURDOUGH":
        return "Tinned House Sourdough"
    if "TINNED SOFT SOURDOUGH" in d or "SOFT SOURDOUGH" in d or "SANDWICH LOAF" in d:
        return "Soft Sourdough"
    if d == "HOUSE SOURDOUGH" or d == "HSE SOURDOUGH":
        return "House Sourdough"
    return None


def main() -> None:
    wb = load_workbook(TRACKER_PATH)
    ws = wb["Tracker"]

    missing: list[tuple[date, str, str, str]] = []
    for wk_mon, base_row in WEEKS:
        for prod_offset, prod in enumerate(PRODUCTS):
            row = base_row + prod_offset
            for d_off in range(7):
                day = wk_mon + timedelta(days=d_off)
                rev_cell = f"{DAY_REV_COLS[d_off]}{row}"
                qty_cell = f"{DAY_QTY_COLS[d_off]}{row}"
                rev_v = ws[rev_cell].value
                qty_v = ws[qty_cell].value
                if rev_v is None and qty_v is None:
                    missing.append((day, prod, rev_cell, qty_cell))

    print(f"Blank product-days found: {len(missing)}")
    needed_dates = sorted({m[0] for m in missing})
    print(f"Distinct dates to fetch from Drive: {len(needed_dates)}")
    if needed_dates:
        print(f"Date span: {needed_dates[0]} to {needed_dates[-1]}")

    service, err = get_service()
    if err:
        raise RuntimeError(err)

    files = list_files_in_folder(service, TMB_SALES_FOLDER_ID, file_pattern=r"\d{8}")
    wanted_set = set(needed_dates)
    wanted_files = []
    for f in files:
        d = extract_date_from_filename(f["name"])
        if d and d.date() in wanted_set:
            wanted_files.append((d.date(), f))

    found_dates = {d for d, _ in wanted_files}
    truly_missing_from_drive = wanted_set - found_dates
    if truly_missing_from_drive:
        print(f"MISSING ON DRIVE: {sorted(truly_missing_from_drive)}")

    daily: dict[date, dict[str, dict[str, float]]] = {}
    for d, f in wanted_files:
        buf = download_file(service, f["id"])
        if buf is None:
            print(f"  download failed: {d}")
            continue
        raw = pd.read_excel(buf)
        raw = raw[raw["Description"].notna()].copy()
        raw["Description"] = raw["Description"].apply(clean_product_name)
        raw["Bucket"] = raw["Description"].apply(categorise)
        raw = raw[raw["Bucket"].notna()]
        raw["Revenue"] = pd.to_numeric(raw["ExtendedNetAmount"], errors="coerce").fillna(0)
        raw["Quantity"] = pd.to_numeric(raw["Quantity"], errors="coerce").fillna(0)
        agg = raw.groupby("Bucket").agg(Rev=("Revenue", "sum"), Qty=("Quantity", "sum"))
        daily[d] = {bucket: {"rev": float(r.Rev), "qty": float(r.Qty)} for bucket, r in agg.iterrows()}
        print(f"  {d}: " + ", ".join(f"{b}={daily[d][b]['rev']:.0f}/{int(daily[d][b]['qty'])}" for b in daily[d]))

    filled = 0
    skipped_no_data = 0
    for day, prod, rev_cell, qty_cell in missing:
        if day not in daily:
            skipped_no_data += 1
            continue
        bucket_data = daily[day].get(prod, {"rev": 0.0, "qty": 0})
        rev = bucket_data["rev"]
        qty = bucket_data["qty"]
        ws[rev_cell] = round(rev, 2) if rev else 0
        ws[qty_cell] = int(qty)
        filled += 1

    print(f"Cells filled: {filled} product-days ({filled*2} cell-values)")
    print(f"Skipped (drive file missing): {skipped_no_data}")

    wb.save(TRACKER_PATH)
    print(f"Saved: {TRACKER_PATH}")


if __name__ == "__main__":
    main()
