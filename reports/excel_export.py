"""Product sales summary Excel export using openpyxl."""

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from config import PDF_HEADER_BG


def generate_excel(filtered_df, min_date, max_date):
    """Build a formatted product sales summary Excel workbook.

    Returns a BytesIO buffer ready for st.download_button.
    """
    if filtered_df.empty:
        buf = BytesIO()
        empty_df = pd.DataFrame(columns=[
            "Product", "Category", "Total Revenue", "Total Quantity",
            "Avg Price", "% of Revenue",
        ])
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            empty_df.to_excel(writer, sheet_name="Product Sales", index=False)
        buf.seek(0)
        return buf

    # --- Aggregate by product ---
    product_summary = (
        filtered_df.groupby(["Description", "Category"], as_index=False)
        .agg(total_revenue=("Revenue", "sum"), total_quantity=("Quantity", "sum"))
    )
    total_rev = product_summary["total_revenue"].sum()
    product_summary["avg_price"] = (
        product_summary["total_revenue"] / product_summary["total_quantity"].replace(0, float("nan"))
    )
    if total_rev:
        product_summary["pct_revenue"] = product_summary["total_revenue"] / total_rev
    else:
        product_summary["pct_revenue"] = 0.0

    # Sort by revenue descending
    product_summary = product_summary.sort_values("total_revenue", ascending=False).reset_index(drop=True)

    # Rename for export
    export_df = product_summary.rename(columns={
        "Description": "Product",
        "Category": "Category",
        "total_revenue": "Total Revenue",
        "total_quantity": "Total Quantity",
        "avg_price": "Avg Price",
        "pct_revenue": "% of Revenue",
    })

    # --- Write to Excel ---
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        export_df.to_excel(writer, sheet_name="Product Sales", index=False)
        ws = writer.sheets["Product Sales"]

        # -- Column widths --
        col_widths = {"A": 30, "B": 18, "C": 14, "D": 14, "E": 12, "F": 14}
        for letter, width in col_widths.items():
            ws.column_dimensions[letter].width = width

        # -- Header styling --
        header_color = PDF_HEADER_BG.lstrip("#")
        header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # -- Number formatting for data rows --
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=6):
            row[2].number_format = '$#,##0.00'    # Total Revenue
            row[3].number_format = '#,##0'         # Total Quantity
            row[4].number_format = '$#,##0.00'    # Avg Price
            row[5].number_format = '0.0%'          # % of Revenue

        # -- Totals row --
        totals_row = ws.max_row + 1
        bold_font = Font(bold=True)
        ws.cell(row=totals_row, column=1, value="TOTAL").font = bold_font
        ws.cell(row=totals_row, column=3, value=total_rev).font = bold_font
        ws.cell(row=totals_row, column=3).number_format = '$#,##0.00'
        total_qty = product_summary["total_quantity"].sum()
        ws.cell(row=totals_row, column=4, value=total_qty).font = bold_font
        ws.cell(row=totals_row, column=4).number_format = '#,##0'

        # -- Freeze header row + auto-filter --
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:F{totals_row - 1}"

    buf.seek(0)
    return buf


def generate_avg_day_excel(filtered_df, selected_day_name, selected_months=None):
    """Build a formatted average-day product performance Excel workbook.

    Returns a BytesIO buffer ready for st.download_button.
    """
    # --- Filter to selected day ---
    df = filtered_df[filtered_df["DayName"] == selected_day_name].copy()

    # --- Optional month filter ---
    if selected_months:
        df = df[df["Date"].dt.month.isin(selected_months)]

    if df.empty:
        # Return an empty workbook with just headers
        buf = BytesIO()
        empty_df = pd.DataFrame(columns=[
            "Product", "Category", "Avg Daily Revenue", "Avg Daily Quantity",
            "Avg Price", "% of Revenue",
        ])
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            empty_df.to_excel(writer, sheet_name="Avg Product Performance", index=False)
        buf.seek(0)
        return buf

    # --- Aggregate: daily totals per product, then average across days ---
    # Use calendar date only (Date column may contain timestamps)
    _date = df["Date"].dt.date
    num_days = _date.nunique()
    daily_product = (
        df.groupby([_date, "Description", "Category"], as_index=False)
        .agg(revenue=("Revenue", "sum"), quantity=("Quantity", "sum"))
    )

    # Divide by total day instances (not just days each product appeared)
    total_product = (
        daily_product.groupby(["Description", "Category"], as_index=False)
        .agg(total_revenue=("revenue", "sum"), total_quantity=("quantity", "sum"))
    )
    total_product["avg_revenue"] = total_product["total_revenue"] / num_days
    total_product["avg_quantity"] = total_product["total_quantity"] / num_days
    avg_product = total_product.drop(columns=["total_revenue", "total_quantity"])

    # Derived columns
    avg_product["avg_price"] = (
        avg_product["avg_revenue"] / avg_product["avg_quantity"].replace(0, float("nan"))
    )
    total_avg_rev = avg_product["avg_revenue"].sum()
    if total_avg_rev:
        avg_product["pct_revenue"] = avg_product["avg_revenue"] / total_avg_rev
    else:
        avg_product["pct_revenue"] = 0.0

    # Sort by avg revenue descending
    avg_product = avg_product.sort_values("avg_revenue", ascending=False).reset_index(drop=True)

    # Rename for export
    export_df = avg_product.rename(columns={
        "Description": "Product",
        "Category": "Category",
        "avg_revenue": "Avg Daily Revenue",
        "avg_quantity": "Avg Daily Quantity",
        "avg_price": "Avg Price",
        "pct_revenue": "% of Revenue",
    })

    # --- Sample size context ---
    num_instances = num_days
    first_date = df["Date"].min().strftime("%d/%m/%Y")
    last_date = df["Date"].max().strftime("%d/%m/%Y")
    context_line = (
        f"Average {selected_day_name} - {num_instances} "
        f"instance{'s' if num_instances != 1 else ''} "
        f"({first_date} to {last_date})"
    )

    # --- Write to Excel ---
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        # Write data starting at row 3 (row 1 = context, row 2 = headers)
        export_df.to_excel(
            writer, sheet_name="Avg Product Performance", index=False, startrow=1,
        )
        ws = writer.sheets["Avg Product Performance"]

        # -- Context row --
        ws.cell(row=1, column=1, value=context_line).font = Font(bold=True, italic=True)

        # -- Column widths --
        col_widths = {"A": 30, "B": 18, "C": 16, "D": 16, "E": 12, "F": 14}
        for letter, width in col_widths.items():
            ws.column_dimensions[letter].width = width

        # -- Header styling (row 2 now) --
        header_color = PDF_HEADER_BG.lstrip("#")
        header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[2]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # -- Number formatting for data rows (row 3 onward) --
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=6):
            row[2].number_format = '$#,##0.00'    # Avg Daily Revenue
            row[3].number_format = '#,##0.0'       # Avg Daily Quantity
            row[4].number_format = '$#,##0.00'    # Avg Price
            row[5].number_format = '0.0%'          # % of Revenue

        # -- Totals row --
        totals_row = ws.max_row + 1
        bold_font = Font(bold=True)
        ws.cell(row=totals_row, column=1, value="TOTAL").font = bold_font
        ws.cell(row=totals_row, column=3, value=total_avg_rev).font = bold_font
        ws.cell(row=totals_row, column=3).number_format = '$#,##0.00'
        total_avg_qty = avg_product["avg_quantity"].sum()
        ws.cell(row=totals_row, column=4, value=total_avg_qty).font = bold_font
        ws.cell(row=totals_row, column=4).number_format = '#,##0.0'

        # -- Freeze below header + auto-filter on data --
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:F{totals_row - 1}"

    buf.seek(0)
    return buf


def generate_baskets_excel(filtered_df):
    """Build a formatted basket-level detail Excel workbook.

    Returns a BytesIO buffer ready for st.download_button.
    """
    if filtered_df.empty:
        buf = BytesIO()
        empty_df = pd.DataFrame(columns=["Basket", "Date", "Items", "Total", "Products"])
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            empty_df.to_excel(writer, sheet_name="Baskets", index=False)
        buf.seek(0)
        return buf

    # --- Aggregate by basket ---
    baskets = (
        filtered_df.groupby("Basket_ID", as_index=False)
        .agg(
            date=("Date", "first"),
            items=("Quantity", "sum"),
            total=("Revenue", "sum"),
            products=("Description", lambda x: ", ".join(sorted(x.unique()))),
        )
    )

    # Sort by date ascending, then total descending
    baskets = baskets.sort_values(["date", "total"], ascending=[True, False]).reset_index(drop=True)

    # Rename for export
    export_df = baskets.rename(columns={
        "Basket_ID": "Basket",
        "date": "Date",
        "items": "Items",
        "total": "Total",
        "products": "Products",
    })

    # --- Write to Excel ---
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        export_df.to_excel(writer, sheet_name="Baskets", index=False)
        ws = writer.sheets["Baskets"]

        # -- Column widths --
        col_widths = {"A": 18, "B": 14, "C": 10, "D": 14, "E": 50}
        for letter, width in col_widths.items():
            ws.column_dimensions[letter].width = width

        # -- Header styling --
        header_color = PDF_HEADER_BG.lstrip("#")
        header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # -- Number formatting for data rows --
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=5):
            row[1].number_format = 'DD/MM/YYYY'   # Date
            row[2].number_format = '#,##0'          # Items
            row[3].number_format = '$#,##0.00'     # Total

        # -- Totals row --
        totals_row = ws.max_row + 1
        bold_font = Font(bold=True)
        ws.cell(row=totals_row, column=1, value="TOTAL").font = bold_font
        total_items = baskets["items"].sum()
        ws.cell(row=totals_row, column=3, value=total_items).font = bold_font
        ws.cell(row=totals_row, column=3).number_format = '#,##0'
        total_revenue = baskets["total"].sum()
        ws.cell(row=totals_row, column=4, value=total_revenue).font = bold_font
        ws.cell(row=totals_row, column=4).number_format = '$#,##0.00'

        # -- Freeze header row + auto-filter --
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:E{totals_row - 1}"

    buf.seek(0)
    return buf
